"""Generate a notional Arrow 3 specification set (Word) for the EDK consistency demo.

All content is fictitious demonstration data: requirement values are notional and every
performance figure is deferred to a (non-existent) classified annex.

baseline/  six baselined documents, uploaded first ("all documents uploaded so far")
incoming/  the incoming Interceptor Segment Specification, with planted inconsistencies
"""
import sys
from pathlib import Path

from docx import Document
from docx.enum.section import WD_SECTION
from docx.enum.table import WD_TABLE_ALIGNMENT
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.shared import Cm, Pt, RGBColor

OUT = Path(sys.argv[1]) if len(sys.argv) > 1 else Path(__file__).resolve().parents[1] / "arrow3-docs"
NAVY = RGBColor(0x0B, 0x2A, 0x4A)
STEEL = RGBColor(0x3E, 0x5C, 0x7A)
GREY = RGBColor(0x5F, 0x6B, 0x7A)
MARKING = "UNCLASSIFIED // FICTITIOUS DEMONSTRATION DATA"
DISCLAIMER = "Notional content generated for an EDK demonstration. Not an official program document."


# ─── Low-level helpers ────────────────────────────────────────────────────────

def shade(cell, hex_fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), hex_fill)
    tc_pr.append(shd)


def field(paragraph, instruction):
    run = paragraph.add_run()
    for kind, text in (("begin", None), (None, instruction), ("separate", None), (None, "1"), ("end", None)):
        if kind:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
            run._r.append(el)
        elif text == "1":
            t = OxmlElement("w:t")
            t.text = "1"
            run._r.append(t)
        else:
            instr = OxmlElement("w:instrText")
            instr.set(qn("xml:space"), "preserve")
            instr.text = f" {text} "
            run._r.append(instr)
    return run


def set_cell(cell, text, *, bold=False, size=9, color=None, align=None):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(0)
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.bold = bold
    run.font.size = Pt(size)
    if color:
        run.font.color.rgb = color


def widths(table, cms):
    # LibreOffice and Word lay out from tblGrid, so set it as well as each cell's width.
    table.autofit = False
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), cms):
        col.set(qn("w:w"), str(Cm(w).twips))
    for row in table.rows:
        for cell, w in zip(row.cells, cms):
            cell.width = Cm(w)


def grid(doc, header, rows, cms, *, header_fill="0B2A4A"):
    table = doc.add_table(rows=1, cols=len(header))
    table.style = "Table Grid"
    table.alignment = WD_TABLE_ALIGNMENT.CENTER
    for cell, text in zip(table.rows[0].cells, header):
        set_cell(cell, text, bold=True, size=8.5, color=RGBColor(0xFF, 0xFF, 0xFF))
        shade(cell, header_fill)
    for i, row in enumerate(rows):
        cells = table.add_row().cells
        for j, (cell, text) in enumerate(zip(cells, row)):
            set_cell(cell, text, bold=(j == 0), size=8.5)
            if i % 2:
                shade(cell, "EEF3F8")
    widths(table, cms)
    doc.add_paragraph().paragraph_format.space_after = Pt(2)
    return table


def header_bar(section, left, right, width_cm, color_left, color_right):
    """Left and right header text in a borderless two-cell table (tab stops render differently in Word and LibreOffice)."""
    header = section.header
    spacer = header.paragraphs[0]
    table = header.add_table(rows=1, cols=2, width=Cm(width_cm))
    tbl_pr = table._tbl.tblPr
    edges = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "nil")
        edges.append(b)
    tbl_pr.append(edges)
    for cell, text, color, align, bold in ((table.rows[0].cells[0], left, color_left, None, True),
                                           (table.rows[0].cells[1], right, color_right, WD_ALIGN_PARAGRAPH.RIGHT, False)):
        cell.width = Cm(width_cm / 2)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if align is not None:
            p.alignment = align
        run = p.add_run(text)
        run.font.size, run.font.color.rgb = Pt(8), color
        run.bold = bold
    spacer._p.addprevious(table._tbl)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.add_run().font.size = Pt(2)


# ─── Document skeleton ────────────────────────────────────────────────────────

def base_document(doc_no, rev):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.2)
    sec.top_margin, sec.bottom_margin = Cm(2.4), Cm(2.2)
    styles = doc.styles
    styles["Normal"].font.name = "Arial"
    styles["Normal"].font.size = Pt(10)
    styles["Normal"].paragraph_format.space_after = Pt(6)
    for name, size, color in (("Heading 1", 14, NAVY), ("Heading 2", 11.5, STEEL), ("Heading 3", 10.5, STEEL)):
        st = styles[name]
        st.font.name, st.font.size, st.font.bold, st.font.color.rgb = "Arial", Pt(size), True, color
        fonts = st.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
        st.paragraph_format.space_before = Pt(12 if name == "Heading 1" else 8)
        st.paragraph_format.space_after = Pt(4)
    sec.different_first_page_header_footer = True
    header_bar(sec, "ARROW 3 WEAPON SYSTEM", f"{doc_no}  Rev {rev}", 16.6, NAVY, GREY)
    foot = sec.footer.paragraphs[0]
    foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = foot.add_run(MARKING + "   ·   Page ")
    r.font.size, r.font.color.rgb = Pt(7.5), GREY
    pr = field(foot, "PAGE")
    pr.font.size = Pt(7.5)
    r = foot.add_run(" of ")
    r.font.size, r.font.color.rgb = Pt(7.5), GREY
    nr = field(foot, "NUMPAGES")
    nr.font.size = Pt(7.5)
    first_foot = sec.first_page_footer.paragraphs[0]
    first_foot.alignment = WD_ALIGN_PARAGRAPH.CENTER
    r = first_foot.add_run(MARKING + "\n" + DISCLAIMER)
    r.font.size, r.font.color.rgb = Pt(7.5), GREY
    return doc


def cover(doc, *, doc_no, rev, title, subtitle, date, status, owner, history, approvals):
    p = doc.add_paragraph()
    p.paragraph_format.space_before = Pt(36)
    r = p.add_run("ARROW 3 WEAPON SYSTEM")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(12), STEEL
    r.font.all_caps = True
    p = doc.add_paragraph()
    r = p.add_run(title)
    r.bold, r.font.size, r.font.color.rgb = True, Pt(26), NAVY
    p = doc.add_paragraph()
    r = p.add_run(subtitle)
    r.font.size, r.font.color.rgb = Pt(13), GREY
    p.paragraph_format.space_after = Pt(28)
    info = doc.add_table(rows=0, cols=2)
    info.style = "Table Grid"
    for k, v in (("Document number", doc_no), ("Revision", rev), ("Date", date), ("Status", status),
                 ("Prepared by", owner), ("Classification", MARKING)):
        cells = info.add_row().cells
        set_cell(cells[0], k, bold=True, size=9, color=NAVY)
        shade(cells[0], "E4ECF4")
        set_cell(cells[1], v, size=9)
    widths(info, [4.2, 12.4])
    doc.add_paragraph().paragraph_format.space_after = Pt(18)
    h = doc.add_paragraph()
    r = h.add_run("Revision history")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(10), NAVY
    grid(doc, ["Rev", "Date", "Description", "CCB ref."], history, [1.4, 2.6, 9.6, 3.0])
    h = doc.add_paragraph()
    r = h.add_run("Approvals")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(10), NAVY
    grid(doc, ["Role", "Organisation", "Signature", "Date"], approvals, [4.6, 5.0, 3.6, 3.4])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def para(doc, text, *, italic=False):
    p = doc.add_paragraph(text)
    if italic:
        for r in p.runs:
            r.italic = True
            r.font.color.rgb = GREY
    return p


def bullets(doc, items):
    for item in items:
        doc.add_paragraph(item, style="List Bullet")


def requirements(doc, rows):
    """rows: (id, text, verification, parent)"""
    grid(doc, ["Req. ID", "Requirement", "Verif.", "Parent / trace"], rows, [2.6, 10.2, 1.4, 2.4])


def applicable(doc, rows):
    grid(doc, ["Document", "Title", "Rev"], rows, [4.0, 11.0, 1.6])


def acronyms(doc, pairs):
    grid(doc, ["Acronym", "Meaning"], pairs, [3.0, 13.6], header_fill="3E5C7A")


ACRONYMS = [("AUR", "All-up round (interceptor in its canister)"), ("BMC", "Battle Management Center"),
            ("CCB", "Configuration Control Board"), ("ECEF", "Earth-Centred, Earth-Fixed"),
            ("FCR", "Fire Control Radar"), ("ICD", "Interface Control Document"),
            ("IFTU", "In-flight target update"), ("LCU", "Launcher Control Unit"),
            ("MTBF", "Mean time between failures"), ("PTL", "Prepare-to-Launch command"),
            ("RT", "Remote terminal (MIL-STD-1553B)"), ("TBR", "To be reviewed")]
VERIFICATION_NOTE = "Verification methods: A = Analysis, I = Inspection, D = Demonstration, T = Test."


# ─── Baseline documents ───────────────────────────────────────────────────────

def srs(path):
    doc_no, rev = "A3-SRS-001", "C"
    doc = base_document(doc_no, rev)
    cover(doc, doc_no=doc_no, rev=rev, title="System Requirements Specification",
          subtitle="Arrow 3 Weapon System (AWS-3) — top-level requirements", date="14 June 2026",
          status="Baselined (SRR action items closed)", owner="Weapon System Engineering Directorate",
          history=[("A", "02 Feb 2025", "Initial release for System Requirements Review", "CCB-25-004"),
                   ("B", "19 Sep 2025", "Timeline allocation added (SRS-TL-002); RAM figures updated", "CCB-25-031"),
                   ("C", "14 Jun 2026", "Time-sync and reference-frame requirements made mandatory for all segments", "CCB-26-017")],
          approvals=[("Chief Systems Engineer", "Weapon System Engineering", "e-signed", "14 Jun 2026"),
                     ("Program Manager", "Program Office", "e-signed", "15 Jun 2026"),
                     ("Quality Assurance", "Product Assurance", "e-signed", "15 Jun 2026")])
    doc.add_heading("1. Scope", 1)
    para(doc, "This System Requirements Specification (SRS) establishes the functional, interface, environmental, "
              "reliability and safety requirements of the Arrow 3 Weapon System (AWS-3), an exo-atmospheric ballistic "
              "missile defence system. It is the parent document of all segment specifications and interface control "
              "documents listed in Section 2.")
    para(doc, "Kill-chain performance values (detection, engagement envelope and probability figures) are contained in the "
              "classified annex A3-SRS-001-CA and are not reproduced in this document.", italic=True)
    doc.add_heading("1.1 System overview", 2)
    para(doc, "AWS-3 comprises five segments, each specified in its own segment specification:")
    bullets(doc, ["Interceptor Segment (INT) — two-stage interceptor with an exo-atmospheric kill vehicle, delivered as an all-up round in a sealed canister.",
                  "Launcher Segment (LCH) — road-mobile launchers, six canister cells each, with a Launcher Control Unit (LCU).",
                  "Fire Control Radar Segment (FCR) — long-range surveillance and fire-control radar.",
                  "Battle Management Center (BMC) — engagement planning, weapon assignment and in-flight target updates.",
                  "Communications Segment (COM) — tactical data network linking all segments."])
    doc.add_heading("2. Applicable documents", 1)
    applicable(doc, [("A3-ICD-110", "Interceptor–Launcher Interface Control Document", "C"),
                     ("A3-ICD-120", "BMC–FCR–Launcher Data Interface Control Document", "C"),
                     ("A3-LCH-SSS-040", "Launcher Segment Specification", "B"),
                     ("A3-FCR-SSS-020", "Fire Control Radar Segment Specification", "B"),
                     ("A3-BMC-SSS-030", "Battle Management Center Specification", "A"),
                     ("A3-ENV-050", "Environmental and RAM Requirements", "B"),
                     ("MIL-STD-1275E", "Characteristics of 28 Volt DC Electrical Systems in Military Vehicles", "E"),
                     ("MIL-STD-461G", "Control of Electromagnetic Interference Characteristics", "G"),
                     ("MIL-STD-810H", "Environmental Engineering Considerations and Laboratory Tests", "H"),
                     ("MIL-STD-882E", "System Safety", "E"),
                     ("NGA.STND.0036", "World Geodetic System 1984 (WGS-84)", "1.0.0")])
    doc.add_heading("3. Requirements", 1)
    para(doc, VERIFICATION_NOTE, italic=True)
    doc.add_heading("3.1 Engagement timeline", 2)
    requirements(doc, [
        ("SRS-TL-001", "The weapon system shall launch the first interceptor within 8.0 s of the engagement order issued by the BMC operator (system reaction time).", "T", "Mission need MN-04"),
        ("SRS-TL-002", "The system reaction time of SRS-TL-001 shall be allocated sequentially: BMC engagement planning (operator order to engagement command at the launcher) ≤ 2.0 s; launcher preparation (engagement command to PTL issued) ≤ 3.0 s; interceptor pre-launch preparation (PTL received to READY reported) ≤ 3.0 s. The launcher issues LCH on READY.", "A", "SRS-TL-001"),
        ("SRS-TL-003", "The system shall support simultaneous engagement of multiple threats; the number of simultaneous engagements is specified in A3-SRS-001-CA.", "D", "Mission need MN-05"),
    ])
    doc.add_heading("3.2 Time and reference frames", 2)
    requirements(doc, [
        ("SRS-TIM-001", "All segments shall maintain a common system time, disciplined to GPS by the BMC master clock, synchronised within ±1 µs of the master clock.", "T", "SRS-TL-001"),
        ("SRS-NAV-002", "All target state vectors and launch parameters exchanged between segments shall be expressed in WGS-84 Earth-Centred, Earth-Fixed (ECEF) coordinates, in metres and metres per second.", "I", "NGA.STND.0036"),
    ])
    doc.add_heading("3.3 Electrical power", 2)
    requirements(doc, [
        ("SRS-PWR-001", "Field equipment shall operate from 28 VDC power with the characteristics of MIL-STD-1275E.", "T", "—"),
        ("SRS-PWR-002", "The launcher shall supply all interceptor pre-launch power through the umbilical defined in A3-ICD-110; the interceptor shall not draw more than the per-canister allocation of A3-ICD-110.", "A", "SRS-PWR-001"),
    ])
    doc.add_heading("3.4 Interfaces", 2)
    requirements(doc, [
        ("SRS-IF-001", "The interceptor–launcher interface shall conform to A3-ICD-110 Rev C.", "I", "—"),
        ("SRS-IF-002", "Data interfaces between BMC, FCR and launchers shall conform to A3-ICD-120 Rev C.", "I", "—"),
        ("SRS-IF-003", "All tactical messages shall be authenticated; encryption per the program cryptographic plan.", "T", "Security CONOPS"),
    ])
    doc.add_heading("3.5 Environment and EMC", 2)
    requirements(doc, [
        ("SRS-ENV-001", "All field-deployed equipment shall meet the environmental requirements of A3-ENV-050 Rev B.", "T", "—"),
        ("SRS-ENV-002", "All field-deployed equipment shall operate at ambient temperatures from −33 °C to +55 °C, with solar loading per MIL-STD-810H Method 505.7.", "T", "SRS-ENV-001"),
        ("SRS-EMC-001", "All segments shall meet MIL-STD-461G, tailored per A3-ENV-050 Rev B.", "T", "—"),
    ])
    doc.add_heading("3.6 Reliability, availability and maintainability", 2)
    requirements(doc, [
        ("SRS-RAM-001", "The weapon system operational availability (Ao) shall be at least 0.95.", "A", "—"),
        ("SRS-RAM-004", "Segment reliability allocations shall be as defined in A3-ENV-050 Rev B, Section 4.", "A", "SRS-RAM-001"),
        ("SRS-RAM-010", "The interceptor shall have a storage life of at least 15 years in its sealed canister, with periodic built-in test only.", "A", "Logistics concept"),
    ])
    doc.add_heading("3.7 Safety", 2)
    requirements(doc, [
        ("SRS-SAF-001", "A system safety program shall be conducted per MIL-STD-882E.", "I", "—"),
        ("SRS-SAF-002", "Rocket motor ignition shall require at least two independent inhibits to be removed, one of which is controlled by the launcher.", "D", "SRS-SAF-001"),
    ])
    doc.add_heading("4. Notes", 1)
    para(doc, "Open items: SRS-TL-003 values are TBR pending the threat assessment update (action SRR-112).")
    acronyms(doc, ACRONYMS)
    doc.save(path)


def icd110(path):
    doc_no, rev = "A3-ICD-110", "C"
    doc = base_document(doc_no, rev)
    cover(doc, doc_no=doc_no, rev=rev, title="Interceptor–Launcher Interface Control Document",
          subtitle="Mechanical, electrical, data and sequencing interface between the all-up round and the launcher",
          date="03 July 2026", status="Baselined", owner="Interface Working Group (INT / LCH)",
          history=[("A", "11 Apr 2025", "Initial release", "CCB-25-012"),
                   ("B", "07 Nov 2025", "Data bus RT addressing and message set Table 4 added", "CCB-25-038"),
                   ("C", "03 Jul 2026", "Pre-launch current limit reduced to 12 A; READY timing aligned with SRS-TL-002", "CCB-26-022")],
          approvals=[("Interface Working Group chair", "Weapon System Engineering", "e-signed", "03 Jul 2026"),
                     ("Interceptor Segment lead", "Interceptor IPT", "e-signed", "04 Jul 2026"),
                     ("Launcher Segment lead", "Launcher IPT", "e-signed", "04 Jul 2026")])
    doc.add_heading("1. Scope", 1)
    para(doc, "This ICD defines the interface between the Arrow 3 all-up round (AUR: interceptor in its canister) and the "
              "launcher canister cell. Both the Interceptor Segment and the Launcher Segment shall comply with it.")
    doc.add_heading("2. Applicable documents", 1)
    applicable(doc, [("A3-SRS-001", "System Requirements Specification", "C"),
                     ("A3-LCH-SSS-040", "Launcher Segment Specification", "B"),
                     ("A3-LCH-DWG-2210", "Canister envelope and interface drawing", "D"),
                     ("MIL-STD-1553B", "Digital Time Division Command/Response Multiplex Data Bus (Notice 4)", "B"),
                     ("MIL-DTL-38999", "Connectors, Electrical, Circular, Series III", "L"),
                     ("MIL-STD-1275E", "28 VDC Electrical Systems in Military Vehicles", "E")])
    doc.add_heading("3. Interface requirements", 1)
    para(doc, VERIFICATION_NOTE, italic=True)
    doc.add_heading("3.1 Mechanical", 2)
    requirements(doc, [
        ("IF-MEC-001", "The canister outer envelope and attachment points shall conform to drawing A3-LCH-DWG-2210 Rev D.", "I", "SRS-IF-001"),
        ("IF-MEC-002", "The mass of the loaded canister (all-up round) shall not exceed 1,350 kg.", "I", "SRS-IF-001"),
        ("IF-MEC-003", "The centre of gravity of the AUR shall lie within the envelope of A3-LCH-DWG-2210 Rev D, sheet 4.", "A", "IF-MEC-001"),
    ])
    doc.add_heading("3.2 Electrical", 2)
    requirements(doc, [
        ("IF-PWR-001", "The launcher shall supply 28 VDC (MIL-STD-1275E) to each canister cell through umbilical connector J1.", "T", "SRS-PWR-002"),
        ("IF-PWR-003", "The current drawn by one AUR from the launcher supply shall not exceed 12 A continuous in any pre-launch state, including IMU run-up and thermal conditioning.", "T", "SRS-PWR-002"),
        ("IF-PWR-004", "Inrush current at power application shall not exceed 20 A for more than 50 ms.", "T", "IF-PWR-003"),
        ("IF-CON-005", "Umbilical connector J1 shall be MIL-DTL-38999 Series III, shell size 25, 61 contacts; pin allocation per Table 3.", "I", "—"),
    ])
    doc.add_heading("3.3 Data", 2)
    requirements(doc, [
        ("IF-DAT-001", "Data between the Launcher Control Unit and the AUR shall be exchanged over a dual-redundant MIL-STD-1553B bus (Bus A / Bus B). The LCU shall be bus controller; the interceptor shall be remote terminal RT 05.", "T", "SRS-IF-001"),
        ("IF-DAT-004", "The message set shall be as defined in Table 4: PTL (Prepare-to-Launch), STS (status), ALN (transfer alignment), TGT (pre-launch target data), LCH (launch) and ABT (abort).", "I", "IF-DAT-001"),
        ("IF-NAV-006", "Pre-launch target data (message TGT) shall be expressed in WGS-84 ECEF coordinates per SRS-NAV-002.", "I", "SRS-NAV-002"),
    ])
    doc.add_heading("3.4 Sequencing and timing", 2)
    requirements(doc, [
        ("IF-SEQ-006", "The interceptor shall report READY (message STS, bit 0) within 3.0 s of receiving PTL.", "T", "SRS-TL-002"),
        ("IF-SEQ-007", "The launcher shall not issue LCH before READY. If READY is not received within 3.5 s of PTL, the LCU shall declare HANG-FIRE and select the next canister.", "T", "SRS-SAF-002"),
        ("IF-TIM-010", "The LCU shall distribute system time to the AUR by 1553B time-tag message plus a 1 PPS discrete; the interceptor shall align its time reference within ±1 µs.", "T", "SRS-TIM-001"),
    ])
    doc.add_heading("4. Notes", 1)
    para(doc, "Table 3 (pin allocation) and Table 4 (message formats) are maintained in the interface database IFDB-A3, export of 01 July 2026.")
    acronyms(doc, ACRONYMS)
    doc.save(path)


def launcher(path):
    doc_no, rev = "A3-LCH-SSS-040", "B"
    doc = base_document(doc_no, rev)
    cover(doc, doc_no=doc_no, rev=rev, title="Launcher Segment Specification",
          subtitle="Road-mobile launcher with six canister cells and Launcher Control Unit", date="21 May 2026",
          status="Baselined", owner="Launcher Integrated Product Team",
          history=[("A", "30 Jun 2025", "Initial release", "CCB-25-019"),
                   ("B", "21 May 2026", "Power distribution and erector rating updated for AUR mass limit", "CCB-26-011")],
          approvals=[("Launcher Segment lead", "Launcher IPT", "e-signed", "21 May 2026"),
                     ("Chief Systems Engineer", "Weapon System Engineering", "e-signed", "22 May 2026")])
    doc.add_heading("1. Scope", 1)
    para(doc, "This specification defines the requirements of the Arrow 3 Launcher Segment: the launcher vehicle, "
              "erector, six canister cells, power distribution unit (PDU) and Launcher Control Unit (LCU).")
    doc.add_heading("2. Applicable documents", 1)
    applicable(doc, [("A3-SRS-001", "System Requirements Specification", "C"),
                     ("A3-ICD-110", "Interceptor–Launcher Interface Control Document", "C"),
                     ("A3-ICD-120", "BMC–FCR–Launcher Data Interface Control Document", "C"),
                     ("A3-ENV-050", "Environmental and RAM Requirements", "B")])
    doc.add_heading("3. Requirements", 1)
    para(doc, VERIFICATION_NOTE, italic=True)
    doc.add_heading("3.1 Mechanical and mobility", 2)
    requirements(doc, [
        ("LCH-MOB-001", "The launcher shall be transportable on public roads without special permits when stowed.", "D", "Logistics concept"),
        ("LCH-MEC-014", "The erector shall be rated for loaded canisters (all-up rounds) of up to 1,350 kg per cell, with a structural safety factor of 1.5.", "A", "IF-MEC-002"),
        ("LCH-MEC-016", "The erector shall raise the canister pack from stowed to launch elevation in not more than 60 s.", "T", "—"),
    ])
    doc.add_heading("3.2 Power", 2)
    requirements(doc, [
        ("LCH-PWR-020", "The launcher shall be capable of autonomous operation from its own generator for at least 24 h.", "D", "—"),
        ("LCH-PWR-021", "The PDU shall supply 28 VDC (MIL-STD-1275E) to each of the six canister cells, limited to 12 A continuous per cell by electronic circuit breakers.", "T", "IF-PWR-003"),
    ])
    doc.add_heading("3.3 Control and data", 2)
    requirements(doc, [
        ("LCH-DAT-010", "The LCU shall act as MIL-STD-1553B bus controller on one dual-redundant bus per canister cell.", "T", "IF-DAT-001"),
        ("LCH-DAT-012", "The LCU shall exchange engagement commands and status with the BMC per A3-ICD-120 Rev C.", "T", "SRS-IF-002"),
        ("LCH-TL-003", "The LCU shall issue PTL to the selected canister within 3.0 s of receiving a valid engagement command, and shall issue LCH within 50 ms of receiving READY.", "T", "SRS-TL-002"),
        ("LCH-TIM-005", "The LCU shall distribute system time to each canister cell within ±1 µs of the BMC master clock.", "T", "SRS-TIM-001"),
    ])
    doc.add_heading("3.4 Environment", 2)
    requirements(doc, [
        ("LCH-ENV-030", "The launcher shall operate at ambient temperatures from −33 °C to +55 °C.", "T", "SRS-ENV-002"),
        ("LCH-ENV-031", "The launcher shall meet the transport shock and vibration requirements of A3-ENV-050 Rev B.", "T", "SRS-ENV-001"),
    ])
    doc.add_heading("3.5 Safety", 2)
    requirements(doc, [
        ("LCH-SAF-040", "The launcher shall provide the launcher-controlled ignition inhibit required by SRS-SAF-002.", "D", "SRS-SAF-002"),
    ])
    acronyms(doc, ACRONYMS)
    doc.save(path)


def fcr(path):
    doc_no, rev = "A3-FCR-SSS-020", "B"
    doc = base_document(doc_no, rev)
    cover(doc, doc_no=doc_no, rev=rev, title="Fire Control Radar Segment Specification",
          subtitle="Surveillance, tracking and fire-control radar for AWS-3", date="09 April 2026",
          status="Baselined", owner="Sensors Integrated Product Team",
          history=[("A", "15 Aug 2025", "Initial release", "CCB-25-027"),
                   ("B", "09 Apr 2026", "Track message rate and time-tag accuracy aligned with SRS Rev C draft", "CCB-26-008")],
          approvals=[("Sensors Segment lead", "Sensors IPT", "e-signed", "09 Apr 2026"),
                     ("Chief Systems Engineer", "Weapon System Engineering", "e-signed", "10 Apr 2026")])
    doc.add_heading("1. Scope", 1)
    para(doc, "This specification defines the requirements of the Fire Control Radar (FCR) Segment of AWS-3.")
    para(doc, "Detection and tracking performance is specified in the classified annex A3-SRS-001-CA.", italic=True)
    doc.add_heading("2. Applicable documents", 1)
    applicable(doc, [("A3-SRS-001", "System Requirements Specification", "C"),
                     ("A3-ICD-120", "BMC–FCR–Launcher Data Interface Control Document", "C"),
                     ("A3-ENV-050", "Environmental and RAM Requirements", "B")])
    doc.add_heading("3. Requirements", 1)
    para(doc, VERIFICATION_NOTE, italic=True)
    requirements(doc, [
        ("FCR-TRK-010", "The FCR shall maintain tracks on all threats within its coverage as specified in A3-SRS-001-CA.", "T", "SRS-TL-003"),
        ("FCR-TRK-012", "The FCR shall report target track state vectors to the BMC at not less than 10 Hz per engaged track, in WGS-84 ECEF coordinates.", "T", "SRS-NAV-002"),
        ("FCR-TIM-003", "Track messages shall carry a time tag accurate to ±1 µs of system time.", "T", "SRS-TIM-001"),
        ("FCR-IF-001", "All FCR–BMC data exchange shall conform to A3-ICD-120 Rev C.", "I", "SRS-IF-002"),
        ("FCR-ENV-001", "The FCR shall meet A3-ENV-050 Rev B for field-deployed equipment.", "T", "SRS-ENV-001"),
        ("FCR-EMC-002", "The FCR shall meet MIL-STD-461G, tailored per A3-ENV-050 Rev B.", "T", "SRS-EMC-001"),
        ("FCR-RAM-004", "The FCR mean time between critical failures (MTBCF) shall be at least 500 h.", "A", "SRS-RAM-004"),
    ])
    acronyms(doc, ACRONYMS)
    doc.save(path)


def bmc(path):
    doc_no, rev = "A3-BMC-SSS-030", "A"
    doc = base_document(doc_no, rev)
    cover(doc, doc_no=doc_no, rev=rev, title="Battle Management Center Specification",
          subtitle="Engagement planning, weapon assignment and in-flight target updates", date="27 February 2026",
          status="Baselined", owner="Battle Management Integrated Product Team",
          history=[("A", "27 Feb 2026", "Initial release", "CCB-26-003")],
          approvals=[("BMC Segment lead", "Battle Management IPT", "e-signed", "27 Feb 2026"),
                     ("Chief Systems Engineer", "Weapon System Engineering", "e-signed", "28 Feb 2026")])
    doc.add_heading("1. Scope", 1)
    para(doc, "This specification defines the requirements of the Battle Management Center (BMC) of AWS-3.")
    doc.add_heading("2. Applicable documents", 1)
    applicable(doc, [("A3-SRS-001", "System Requirements Specification", "C"),
                     ("A3-ICD-120", "BMC–FCR–Launcher Data Interface Control Document", "C")])
    doc.add_heading("3. Requirements", 1)
    para(doc, VERIFICATION_NOTE, italic=True)
    requirements(doc, [
        ("BMC-ENG-005", "The BMC shall complete engagement planning and transmit the engagement command to the selected launcher within 2.0 s of the operator engagement order.", "T", "SRS-TL-002"),
        ("BMC-ENG-020", "The BMC shall transmit in-flight target updates (IFTU) to each interceptor in flight at not less than 4 Hz, in WGS-84 ECEF coordinates, per A3-ICD-120 Rev C.", "T", "SRS-NAV-002"),
        ("BMC-TIM-002", "The BMC master clock shall be GPS-disciplined and shall distribute system time to all segments with an accuracy of ±1 µs.", "T", "SRS-TIM-001"),
        ("BMC-NAV-004", "The BMC shall perform all engagement computations and exchange all state vectors in WGS-84 ECEF.", "I", "SRS-NAV-002"),
        ("BMC-CYB-010", "All BMC tactical messages shall be authenticated and encrypted per the program cryptographic plan.", "T", "SRS-IF-003"),
        ("BMC-RAM-012", "The BMC shall achieve an inherent availability of at least 0.99 with hot-standby redundancy.", "A", "SRS-RAM-001"),
    ])
    acronyms(doc, ACRONYMS)
    doc.save(path)


def env(path):
    doc_no, rev = "A3-ENV-050", "B"
    doc = base_document(doc_no, rev)
    cover(doc, doc_no=doc_no, rev=rev, title="Environmental and RAM Requirements",
          subtitle="Climatic, dynamic, EMC and reliability requirements for all AWS-3 segments", date="30 January 2026",
          status="Baselined", owner="Specialty Engineering",
          history=[("A", "22 May 2025", "Initial release", "CCB-25-015"),
                   ("B", "30 Jan 2026", "Low operating temperature extended to −33 °C; interceptor MTBF raised to 2,000 h", "CCB-26-001")],
          approvals=[("Specialty Engineering lead", "Weapon System Engineering", "e-signed", "30 Jan 2026"),
                     ("Chief Systems Engineer", "Weapon System Engineering", "e-signed", "02 Feb 2026")])
    doc.add_heading("1. Scope", 1)
    para(doc, "This document defines the natural and induced environments, EMC requirements and reliability allocations "
              "that apply to every field-deployed item of AWS-3.")
    doc.add_heading("2. Applicable documents", 1)
    applicable(doc, [("A3-SRS-001", "System Requirements Specification", "C"),
                     ("MIL-STD-810H", "Environmental Engineering Considerations and Laboratory Tests", "H"),
                     ("MIL-STD-461G", "Control of Electromagnetic Interference Characteristics", "G")])
    doc.add_heading("3. Environmental requirements", 1)
    para(doc, VERIFICATION_NOTE, italic=True)
    requirements(doc, [
        ("ENV-TMP-001", "All field-deployed equipment, including the all-up round in its canister, shall operate at ambient temperatures from −33 °C to +55 °C.", "T", "SRS-ENV-002"),
        ("ENV-TMP-002", "Non-operating (storage and transport) temperature range shall be −40 °C to +71 °C.", "T", "SRS-ENV-001"),
        ("ENV-HUM-003", "Equipment shall withstand humidity per MIL-STD-810H Method 507.6, Procedure II.", "T", "SRS-ENV-001"),
        ("ENV-SND-004", "Equipment shall withstand blowing sand and dust per MIL-STD-810H Method 510.7.", "T", "SRS-ENV-001"),
        ("ENV-VIB-005", "Transport vibration shall be per MIL-STD-810H Method 514.8, Category 4 (truck transportation).", "T", "SRS-ENV-001"),
        ("ENV-EMC-010", "All segments shall meet MIL-STD-461G requirements CE102, CS101, CS114, RE102 and RS103 as tailored in Annex B.", "T", "SRS-EMC-001"),
    ])
    doc.add_heading("4. Reliability allocations", 1)
    requirements(doc, [
        ("RAM-004", "Interceptor (all-up round) captive-carry MTBF, powered in the canister on the launcher, shall be at least 2,000 h.", "A", "SRS-RAM-004"),
        ("RAM-005", "Launcher MTBF shall be at least 1,200 h.", "A", "SRS-RAM-004"),
        ("RAM-006", "FCR MTBCF shall be at least 500 h.", "A", "SRS-RAM-004"),
        ("RAM-010", "Interceptor storage life in the sealed canister shall be at least 15 years.", "A", "SRS-RAM-010"),
    ])
    acronyms(doc, ACRONYMS)
    doc.save(path)


# ─── The incoming document (planted inconsistencies marked ✱ in comments) ─────

def interceptor(path):
    doc_no, rev = "A3-INT-SSS-010", "A"
    doc = base_document(doc_no, rev)
    cover(doc, doc_no=doc_no, rev=rev, title="Interceptor Segment Specification",
          subtitle="Two-stage interceptor and exo-atmospheric kill vehicle — draft for Preliminary Design Review",
          date="20 October 2026", status="Draft for PDR — submitted for consistency check",
          owner="Interceptor Integrated Product Team",
          history=[("A", "20 Oct 2026", "First issue for Preliminary Design Review", "—")],
          approvals=[("Interceptor Segment lead", "Interceptor IPT", "e-signed", "20 Oct 2026"),
                     ("Chief Systems Engineer", "Weapon System Engineering", "pending", "—")])
    doc.add_heading("1. Scope", 1)
    para(doc, "This specification defines the requirements of the Arrow 3 Interceptor Segment: the two-stage interceptor, "
              "the exo-atmospheric kill vehicle and the canister, delivered together as the all-up round (AUR).")
    para(doc, "Kinematic and lethality performance is specified in the classified annex A3-SRS-001-CA.", italic=True)
    doc.add_heading("2. Applicable documents", 1)
    applicable(doc, [("A3-SRS-001", "System Requirements Specification", "C"),
                     ("A3-ICD-110", "Interceptor–Launcher Interface Control Document", "B"),        # ✱ baseline is Rev C
                     ("A3-ICD-120", "BMC–FCR–Launcher Data Interface Control Document", "A"),       # ✱ baseline is Rev C
                     ("A3-ENV-050", "Environmental and RAM Requirements", "B"),
                     ("MIL-STD-461F", "Control of Electromagnetic Interference Characteristics", "F"),  # ✱ baseline is 461G
                     ("MIL-STD-1275E", "28 VDC Electrical Systems in Military Vehicles", "E")])
    doc.add_heading("3. Requirements", 1)
    para(doc, VERIFICATION_NOTE, italic=True)
    doc.add_heading("3.1 General and physical", 2)
    requirements(doc, [
        ("INT-GEN-001", "The interceptor shall be a two-stage, solid-propellant vehicle carrying an exo-atmospheric hit-to-kill kill vehicle.", "I", "A3-SRS-001 §1.1"),
        ("INT-PHY-002", "The mass of the all-up round, including the canister, shall not exceed 1,450 kg.", "I", "A3-ICD-110"),            # ✱ 1,350 kg
        ("INT-PHY-003", "The canister envelope shall conform to drawing A3-LCH-DWG-2210 Rev D.", "I", "IF-MEC-001"),
    ])
    doc.add_heading("3.2 Electrical", 2)
    requirements(doc, [
        ("INT-PWR-010", "The interceptor shall operate from launcher-supplied 28 VDC power with the characteristics of MIL-STD-1275E.", "T", "SRS-PWR-001"),
        ("INT-PWR-012", "During pre-launch (IMU run-up and thermal conditioning) the interceptor shall draw not more than 18 A continuous from the launcher supply.", "T", "SRS-PWR-002"),  # ✱ 12 A
        ("INT-PWR-014", "Internal thermal batteries shall be activated on receipt of the LCH command and shall power the interceptor for the whole flight.", "T", "—"),
    ])
    doc.add_heading("3.3 Data and timing", 2)
    requirements(doc, [
        ("INT-DAT-020", "The interceptor shall exchange pre-launch data with the launcher over a dual-redundant Gigabit Ethernet (1000BASE-T) link, the interceptor acting as server.", "T", "A3-ICD-110"),  # ✱ MIL-STD-1553B
        ("INT-DAT-022", "In-flight target update messages shall be decoded per A3-ICD-120 Rev A.", "I", "A3-ICD-120"),   # ✱ Rev C
        ("INT-TIM-031", "The interceptor shall report READY within 6.0 s of receiving the Prepare-to-Launch (PTL) command.", "T", "SRS-TL-002"),  # ✱ 3.0 s
        ("INT-TIM-035", "The interceptor time reference shall be synchronised to system time within ±10 µs.", "T", "SRS-TIM-001"),  # ✱ ±1 µs
    ])
    doc.add_heading("3.4 Navigation and guidance", 2)
    requirements(doc, [
        ("INT-NAV-040", "Pre-launch target data and in-flight target updates shall be received and processed in Earth-Centred Inertial coordinates (ECI, J2000 epoch).", "I", "SRS-NAV-002"),  # ✱ WGS-84 ECEF
        ("INT-NAV-042", "The inertial measurement unit shall be aligned before launch by transfer alignment from the launcher (message ALN).", "T", "IF-DAT-004"),
    ])
    doc.add_heading("3.5 Environment and EMC", 2)
    requirements(doc, [
        ("INT-ENV-050", "The all-up round shall operate at ambient temperatures from −20 °C to +50 °C.", "T", "SRS-ENV-002"),  # ✱ −33..+55 °C
        ("INT-ENV-052", "Non-operating (storage and transport) temperature range shall be −40 °C to +71 °C.", "T", "ENV-TMP-002"),
        ("INT-EMC-060", "The interceptor shall meet MIL-STD-461F requirements CE102, CS114, RE102 and RS103.", "T", "SRS-EMC-001"),  # ✱ 461G
    ])
    doc.add_heading("3.6 Reliability and safety", 2)
    requirements(doc, [
        ("INT-RAM-070", "The captive-carry MTBF of the all-up round, powered on the launcher, shall be at least 1,500 h.", "A", "RAM-004"),  # ✱ 2,000 h
        ("INT-RAM-072", "The interceptor shall have a storage life of at least 10 years in its sealed canister.", "A", "SRS-RAM-010"),  # ✱ 15 years
        ("INT-SAF-080", "Rocket motor ignition shall require two independent inhibits to be removed, one controlled by the launcher.", "D", "SRS-SAF-002"),
        ("INT-CYB-090", "All messages received by the interceptor shall be authenticated before use.", "T", "SRS-IF-003"),
    ])
    doc.add_heading("4. Notes", 1)
    para(doc, "This draft is submitted for the PDR data package. Segment-level changes since the concept phase: Ethernet "
              "replaces the legacy data bus to support faster alignment transfer; the guidance team prefers inertial "
              "coordinates for in-flight updates.")
    acronyms(doc, ACRONYMS)
    doc.save(path)


if __name__ == "__main__":
    (OUT / "baseline").mkdir(parents=True, exist_ok=True)
    (OUT / "incoming").mkdir(parents=True, exist_ok=True)
    srs(OUT / "baseline/A3-SRS-001_RevC_System_Requirements_Specification.docx")
    icd110(OUT / "baseline/A3-ICD-110_RevC_Interceptor-Launcher_ICD.docx")
    launcher(OUT / "baseline/A3-LCH-SSS-040_RevB_Launcher_Segment_Specification.docx")
    fcr(OUT / "baseline/A3-FCR-SSS-020_RevB_Fire_Control_Radar_Specification.docx")
    bmc(OUT / "baseline/A3-BMC-SSS-030_RevA_Battle_Management_Center_Specification.docx")
    env(OUT / "baseline/A3-ENV-050_RevB_Environmental_and_RAM_Requirements.docx")
    interceptor(OUT / "incoming/A3-INT-SSS-010_RevA_Interceptor_Segment_Specification.docx")
    print("written to", OUT)
