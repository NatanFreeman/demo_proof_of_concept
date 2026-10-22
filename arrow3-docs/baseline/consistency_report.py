#!/usr/bin/env python3
"""Arrow 3 program report kit: turn consistency findings into a Word report.

    python3 consistency_report.py FINDINGS... --index INDEX.md --source NEW.docx -o REPORT.docx

FINDINGS is the Markdown/JSON file written by the "Find inconsistencies" step: a JSON array
of objects with id, severity, category, title, new_req, new_text, baseline_req, baseline_doc,
baseline_text, impact and recommendation. Missing keys are tolerated; a file without any
parsable finding is an error (the report would claim the document is consistent).
"""
import argparse
import datetime as dt
import json
import re
import sys
import tempfile
from collections import Counter, defaultdict
from pathlib import Path

import matplotlib
matplotlib.use("Agg")
import matplotlib.pyplot as plt  # noqa: E402
from docx import Document  # noqa: E402
from docx.enum.table import WD_CELL_VERTICAL_ALIGNMENT, WD_TABLE_ALIGNMENT  # noqa: E402
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_BREAK  # noqa: E402
from docx.oxml import OxmlElement  # noqa: E402
from docx.oxml.ns import qn  # noqa: E402
from docx.shared import Cm, Pt, RGBColor  # noqa: E402

NAVY, INK, GREY = "0B2A4A", "1F2937", "6B7280"
SEVERITIES = ("Critical", "Major", "Minor")
SEV_FILL = {"Critical": "B91C1C", "Major": "EA580C", "Minor": "CA8A04"}
SEV_TINT = {"Critical": "FDE8E8", "Major": "FFEDD5", "Minor": "FEF9C3"}
MARKING = "UNCLASSIFIED // FICTITIOUS DEMONSTRATION DATA"
DOC_NO = re.compile(r"A3-[A-Z]{2,4}(?:-[A-Z]{3})?-\d{3}")
REQ_ID = re.compile(r"\b(?:[A-Z]{2,4}-[A-Z]{2,4}-\d{3}|RAM-\d{3})\b")
OWNERS = {"power": "Interface Working Group (INT/LCH)", "electrical": "Interface Working Group (INT/LCH)",
          "data": "Interface Working Group (INT/LCH)", "interface": "Interface Working Group (INT/LCH)",
          "mechanical": "Interface Working Group (INT/LCH)", "mass": "Interface Working Group (INT/LCH)",
          "timing": "Weapon System Engineering", "timeline": "Weapon System Engineering",
          "time": "Weapon System Engineering", "navigation": "Weapon System Engineering",
          "coordinate": "Weapon System Engineering", "environment": "Specialty Engineering",
          "emc": "Specialty Engineering", "reliability": "Specialty Engineering", "ram": "Specialty Engineering",
          "document": "Configuration Management", "reference": "Configuration Management"}


# ─── Input ────────────────────────────────────────────────────────────────────

def parse_findings(text: str) -> list[dict]:
    text = re.sub(r"```(?:json)?", "", text)
    start, end = text.find("["), text.rfind("]")
    if start != -1 and end > start:
        try:
            data = json.loads(text[start:end + 1])
            if isinstance(data, list):
                return [d for d in data if isinstance(d, dict)]
        except json.JSONDecodeError:
            pass
    found, depth, begin = [], 0, None
    for i, ch in enumerate(text):  # salvage individual objects from a damaged array
        if ch == "{":
            depth, begin = depth + 1, i if depth == 0 else begin
        elif ch == "}" and depth:
            depth -= 1
            if depth == 0:
                try:
                    found.append(json.loads(text[begin:i + 1]))
                except json.JSONDecodeError:
                    pass
    return [d for d in found if isinstance(d, dict)]


def pick(d: dict, *keys, default=""):
    lowered = {re.sub(r"[^a-z]", "", k.lower()): v for k, v in d.items()}
    for key in keys:
        value = lowered.get(re.sub(r"[^a-z]", "", key.lower()))
        if value not in (None, "", []):
            return " ".join(map(str, value)) if isinstance(value, list) else str(value).strip()
    return default


def severity(value: str) -> str:
    v = value.lower()
    if any(w in v for w in ("crit", "high", "blocker", "sev1")):
        return "Critical"
    if any(w in v for w in ("maj", "med", "sev2")):
        return "Major"
    return "Minor"


def normalise(raw: list[dict]) -> list[dict]:
    findings = []
    for d in raw:
        f = {"severity": severity(pick(d, "severity", "sev", "priority")),
             "category": pick(d, "category", "type", "area", default="General").title(),
             "title": pick(d, "title", "summary", "issue", "description", default="Inconsistency"),
             "new_req": pick(d, "new_req", "new_requirement", "new_id", "new_req_id"),
             "new_text": pick(d, "new_text", "new_quote", "new_requirement_text"),
             "baseline_req": pick(d, "baseline_req", "baseline_id", "existing_req", "baseline_req_id"),
             "baseline_doc": pick(d, "baseline_doc", "baseline_document", "document", "baseline_file"),
             "baseline_text": pick(d, "baseline_text", "baseline_quote", "existing_text"),
             "impact": pick(d, "impact", "consequence", "risk"),
             "recommendation": pick(d, "recommendation", "resolution", "action", "fix")}
        if f["new_req"] or f["baseline_req"] or f["title"] != "Inconsistency":
            findings.append(f)
    findings.sort(key=lambda f: (SEVERITIES.index(f["severity"]), f["new_req"]))
    for number, f in enumerate(findings, start=1):
        f["id"] = f"F-{number:02d}"
        f["doc"] = short_doc(f["baseline_doc"]) or "Baseline"
    return findings


def short_doc(text: str) -> str:
    match = DOC_NO.search(text or "")
    return match.group(0) if match else Path(text).stem[:28] if text else ""


def describe(path: str) -> dict:
    stem = Path(path).stem
    match = re.match(r"(A3-[A-Z0-9-]+?)_Rev([A-Z0-9]+)_(.+)", stem)
    if match:
        return {"no": match.group(1), "rev": match.group(2), "title": match.group(3).replace("_", " ").replace("-", "–")}
    return {"no": short_doc(stem) or stem, "rev": "—", "title": stem.replace("_", " ")}


def parse_index(path: Path) -> list[dict]:
    if not path.is_file():
        return []
    sections, current = [], None
    for line in path.read_text(errors="replace").splitlines():
        if line.startswith("## "):
            current = {"path": line[3:].strip().strip("`"), "ids": set()}
            sections.append(current)
        elif current is not None and line.lstrip().startswith("|"):
            # Only requirement table rows count (a header line like "A3-BMC-SSS-030, Rev A" would match the ID
            # pattern); a row's own ID comes first, later IDs are its parent/trace references.
            first = REQ_ID.search(line)
            if first:
                current["ids"].add(first.group(0))
    return [{**describe(s["path"]), "path": s["path"], "requirements": len(s["ids"])} for s in sections]


# ─── Word helpers ─────────────────────────────────────────────────────────────

def shade(cell, fill):
    tc_pr = cell._tc.get_or_add_tcPr()
    for old in tc_pr.findall(qn("w:shd")):
        tc_pr.remove(old)
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:color"), "auto")
    shd.set(qn("w:fill"), fill)
    tc_pr.append(shd)


def borders(table, color="D1D5DB", size=4):
    tbl_pr = table._tbl.tblPr
    el = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "single" if size else "nil")
        b.set(qn("w:sz"), str(size))
        b.set(qn("w:color"), color)
        el.append(b)
    tbl_pr.append(el)


def widths(table, cms):
    table.autofit = False
    for col, w in zip(table._tbl.tblGrid.findall(qn("w:gridCol")), cms):
        col.set(qn("w:w"), str(Cm(w).twips))
    for row in table.rows:
        for cell, w in zip(row.cells, cms):
            cell.width = Cm(w)


def write(cell, text, *, size=9, bold=False, color=INK, align=None, italic=False, space=0):
    cell.text = ""
    p = cell.paragraphs[0]
    p.paragraph_format.space_after = Pt(space)
    if align:
        p.alignment = align
    run = p.add_run(text)
    run.font.size, run.bold, run.italic = Pt(size), bold, italic
    run.font.color.rgb = RGBColor.from_string(color)
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
    return p


def heading(doc, text, level=1):
    h = doc.add_heading(text, level)
    return h


def field(paragraph, instruction, size):
    run = paragraph.add_run()
    run.font.size = Pt(size)
    for kind in ("begin", "instr", "separate", "text", "end"):
        if kind == "instr":
            el = OxmlElement("w:instrText")
            el.set(qn("xml:space"), "preserve")
            el.text = f" {instruction} "
        elif kind == "text":
            el = OxmlElement("w:t")
            el.text = "1"
        else:
            el = OxmlElement("w:fldChar")
            el.set(qn("w:fldCharType"), kind)
        run._r.append(el)


def setup(doc, subject):
    sec = doc.sections[0]
    sec.page_width, sec.page_height = Cm(21), Cm(29.7)
    sec.left_margin = sec.right_margin = Cm(2.0)
    sec.top_margin, sec.bottom_margin = Cm(2.0), Cm(1.8)
    st = doc.styles
    # Requirement IDs (INT-NAV-040, ...) are not words: mark all text "do not check spelling" so Word and
    # LibreOffice show no red squiggles.
    # Word honours w:noProof; LibreOffice honours the "no linguistic content" language (zxx).
    defaults = st.element.find(qn("w:docDefaults")).find(qn("w:rPrDefault")).find(qn("w:rPr"))
    defaults.append(OxmlElement("w:noProof"))
    lang = defaults.find(qn("w:lang"))
    if lang is None:
        lang = OxmlElement("w:lang")
        defaults.append(lang)
    for attr in ("w:val", "w:eastAsia", "w:bidi"):
        lang.set(qn(attr), "zxx")
    st["Normal"].font.name, st["Normal"].font.size = "Arial", Pt(10)
    st["Normal"].font.color.rgb = RGBColor.from_string(INK)
    st["Normal"].paragraph_format.space_after = Pt(5)
    for name, size, color in (("Heading 1", 16, NAVY), ("Heading 2", 12.5, "1E4E79"), ("Heading 3", 11, "1E4E79")):
        s = st[name]
        s.font.name, s.font.size, s.font.bold = "Arial", Pt(size), True
        s.font.color.rgb = RGBColor.from_string(color)
        fonts = s.element.rPr.rFonts
        for attr in ("w:asciiTheme", "w:hAnsiTheme", "w:eastAsiaTheme", "w:cstheme"):
            fonts.attrib.pop(qn(attr), None)
        s.paragraph_format.space_before, s.paragraph_format.space_after = Pt(14 if name == "Heading 1" else 10), Pt(5)
    sec.different_first_page_header_footer = True
    header_bar(sec, "ARROW 3  ·  REQUIREMENTS CONSISTENCY REPORT", subject, 17,
               RGBColor.from_string(NAVY), RGBColor.from_string(GREY))
    for footer in (sec.footer, sec.first_page_footer):
        p = footer.paragraphs[0]
        p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(MARKING + "   ·   Page ")
        r.font.size, r.font.color.rgb = Pt(7.5), RGBColor.from_string(GREY)
        field(p, "PAGE", 7.5)
        r = p.add_run(" of ")
        r.font.size, r.font.color.rgb = Pt(7.5), RGBColor.from_string(GREY)
        field(p, "NUMPAGES", 7.5)


def header_bar(section, left, right, width_cm, color_left, color_right):
    """Left and right header text in a borderless two-cell table (tab stops render differently in Word and LibreOffice)."""
    header = section.header
    spacer = header.paragraphs[0]
    table = header.add_table(rows=1, cols=2, width=Cm(width_cm))
    tbl_pr = table._tbl.tblPr
    edges = OxmlElement("w:tblBorders")
    for edge in ("top", "left", "bottom", "right", "insideH", "insideV"):
        b = OxmlElement(f"w:{edge}")
        b.set(qn("w:val"), "nil")
        edges.append(b)
    tbl_pr.append(edges)
    for cell, text, color, align, bold in ((table.rows[0].cells[0], left, color_left, None, True),
                                           (table.rows[0].cells[1], right, color_right, WD_ALIGN_PARAGRAPH.RIGHT, False)):
        cell.width = Cm(width_cm / 2)
        p = cell.paragraphs[0]
        p.paragraph_format.space_after = Pt(0)
        if align is not None:
            p.alignment = align
        run = p.add_run(text)
        run.font.size, run.font.color.rgb = Pt(8), color
        run.bold = bold
    spacer._p.addprevious(table._tbl)
    spacer.paragraph_format.space_after = Pt(0)
    spacer.add_run().font.size = Pt(2)


# ─── Charts ───────────────────────────────────────────────────────────────────

def charts(findings, folder: Path) -> Path:
    counts = Counter(f["severity"] for f in findings)
    per_doc = defaultdict(Counter)
    for f in findings:
        per_doc[f["doc"]][f["severity"]] += 1
    plt.rcParams.update({"font.family": "DejaVu Sans", "font.size": 10})
    fig = plt.figure(figsize=(10, 3.6), dpi=200)
    ax1 = fig.add_axes([0.0, 0.05, 0.32, 0.9])
    values = [counts.get(s, 0) for s in SEVERITIES]
    shown = [(v, s) for v, s in zip(values, SEVERITIES) if v]
    ax1.pie([v for v, _ in shown], colors=["#" + SEV_FILL[s] for _, s in shown], startangle=90, counterclock=False,
            wedgeprops={"width": 0.36, "edgecolor": "white", "linewidth": 2})
    ax1.text(0, 0.08, str(len(findings)), ha="center", va="center", fontsize=30, fontweight="bold", color="#" + NAVY)
    ax1.text(0, -0.22, "findings", ha="center", va="center", fontsize=10, color="#" + GREY)
    ax1.set_title("By severity", fontsize=11, fontweight="bold", color="#" + NAVY, pad=2)
    ax2 = fig.add_axes([0.47, 0.16, 0.5, 0.72])
    docs = sorted(per_doc, key=lambda d: (-sum(per_doc[d].values()), d))
    left = [0] * len(docs)
    for s in SEVERITIES:
        widths_ = [per_doc[d].get(s, 0) for d in docs]
        ax2.barh(docs, widths_, left=left, color="#" + SEV_FILL[s], label=s, height=0.62, edgecolor="white")
        left = [a + b for a, b in zip(left, widths_)]
    ax2.invert_yaxis()
    ax2.set_title("Conflicting baseline document", fontsize=11, fontweight="bold", color="#" + NAVY, loc="left")
    ax2.xaxis.get_major_locator().set_params(integer=True)
    for spine in ("top", "right"):
        ax2.spines[spine].set_visible(False)
    ax2.legend(frameon=False, ncol=3, loc="lower right", fontsize=9, bbox_to_anchor=(1.0, -0.32))
    path = folder / "summary.png"
    fig.savefig(path, bbox_inches="tight", facecolor="white")
    plt.close(fig)
    return path


# ─── Report ───────────────────────────────────────────────────────────────────

def cover(doc, new, baseline, findings, generated):
    counts = Counter(f["severity"] for f in findings)
    banner = doc.add_table(rows=1, cols=1)
    borders(banner, size=0)
    cell = banner.rows[0].cells[0]
    shade(cell, NAVY)
    p = write(cell, "ARROW 3 WEAPON SYSTEM", size=10, bold=True, color="9FB7D3", space=2)
    p.paragraph_format.space_before = Pt(14)
    p = cell.add_paragraph()
    r = p.add_run("Requirements Consistency Report")
    r.bold, r.font.size, r.font.color.rgb = True, Pt(26), RGBColor(0xFF, 0xFF, 0xFF)
    p = cell.add_paragraph()
    p.paragraph_format.space_after = Pt(16)
    r = p.add_run(f"{new['no']} Rev {new['rev']}  ·  {new['title']}  vs.  the program baseline")
    r.font.size, r.font.color.rgb = Pt(11), RGBColor(0xDB, 0xE7, 0xF3)
    for par in cell.paragraphs:
        par.paragraph_format.left_indent = Cm(0.5)
    widths(banner, [17])
    doc.add_paragraph().paragraph_format.space_after = Pt(6)

    tiles = doc.add_table(rows=2, cols=4)
    borders(tiles, color="FFFFFF", size=12)
    for i, (label, value, fill) in enumerate((("TOTAL FINDINGS", len(findings), NAVY),
                                              ("CRITICAL", counts.get("Critical", 0), SEV_FILL["Critical"]),
                                              ("MAJOR", counts.get("Major", 0), SEV_FILL["Major"]),
                                              ("MINOR", counts.get("Minor", 0), SEV_FILL["Minor"]))):
        top, bottom = tiles.rows[0].cells[i], tiles.rows[1].cells[i]
        shade(top, fill)
        shade(bottom, fill)
        write(top, str(value), size=34, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER).paragraph_format.space_before = Pt(8)
        write(bottom, label, size=8.5, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER, space=8)
    widths(tiles, [4.25] * 4)
    doc.add_paragraph().paragraph_format.space_after = Pt(4)

    critical, major = counts.get("Critical", 0), counts.get("Major", 0)
    if critical:
        verdict, fill, tint = (f"NOT READY FOR BASELINE — {critical} critical and {major} major inconsistencies must be "
                               "resolved before the document can enter the PDR baseline."), SEV_FILL["Critical"], SEV_TINT["Critical"]
    elif major:
        verdict, fill, tint = f"CONDITIONAL — resolve {major} major inconsistencies before PDR.", SEV_FILL["Major"], SEV_TINT["Major"]
    else:
        verdict, fill, tint = "CONSISTENT — no blocking inconsistencies with the baseline.", "15803D", "DCFCE7"
    box = doc.add_table(rows=1, cols=2)
    borders(box, size=0)
    shade(box.rows[0].cells[0], fill)
    shade(box.rows[0].cells[1], tint)
    write(box.rows[0].cells[0], "VERDICT", size=9, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    write(box.rows[0].cells[1], verdict, size=10, bold=True, color=fill, space=0).paragraph_format.space_before = Pt(6)
    box.rows[0].cells[1].paragraphs[0].paragraph_format.space_after = Pt(6)
    widths(box, [2.6, 14.4])
    doc.add_paragraph().paragraph_format.space_after = Pt(10)

    requirements = sum(d["requirements"] for d in baseline)
    info = doc.add_table(rows=0, cols=2)
    borders(info)
    for key, value in (("Document under review", f"{new['no']} Rev {new['rev']} — {new['title']}"),
                       ("Baseline checked", f"{len(baseline)} documents uploaded so far · {requirements} requirements indexed"),
                       ("Generated", generated.strftime("%d %B %Y, %H:%M")),
                       ("Produced by", "EDK custom agent “Find Inconsistencies in Requirements” (two Hermes agents) — automated analysis, "
                                       "to be confirmed by the Interface Working Group"),
                       ("Classification", MARKING)):
        cells = info.add_row().cells
        shade(cells[0], "EEF2F7")
        write(cells[0], key, size=8.5, bold=True, color=NAVY, space=3)
        write(cells[1], value, size=8.5, space=3)
    widths(info, [4.4, 12.6])
    crit = [f for f in findings if f["severity"] == "Critical"] or [f for f in findings if f["severity"] == "Major"]
    if crit:
        p = doc.add_paragraph()
        p.paragraph_format.space_before = Pt(16)
        r = p.add_run("Blocking findings" if crit[0]["severity"] == "Critical" else "Most serious findings")
        r.bold, r.font.size, r.font.color.rgb = True, Pt(11), RGBColor.from_string(NAVY)
        table = doc.add_table(rows=0, cols=3)
        borders(table, color="E5E7EB")
        for f in crit:
            cells = table.add_row().cells
            shade(cells[0], SEV_FILL[f["severity"]])
            write(cells[0], f["id"], size=9, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
            write(cells[1], f["title"], size=9, bold=True, space=3)
            write(cells[2], f"{f['new_req']}  ↔  {f['baseline_req']}\n{f['doc']}", size=8, color=GREY)
        widths(table, [1.5, 11.0, 4.5])
    doc.add_paragraph().add_run().add_break(WD_BREAK.PAGE)


def summary(doc, new, baseline, findings, chart):
    heading(doc, "1. Executive summary")
    counts = Counter(f["severity"] for f in findings)
    docs = Counter(f["doc"] for f in findings)
    top_doc, top_n = docs.most_common(1)[0] if docs else ("—", 0)
    doc.add_paragraph(
        f"The incoming {new['title']} ({new['no']} Rev {new['rev']}) was checked requirement by requirement against "
        f"all {len(baseline)} Arrow 3 documents uploaded so far. The check found {len(findings)} inconsistencies: "
        f"{counts.get('Critical', 0)} critical, {counts.get('Major', 0)} major and {counts.get('Minor', 0)} minor. "
        f"Most conflicts are with {top_doc} ({top_n}).")
    doc.add_paragraph("Critical findings break a baselined interface and block the PDR baseline; major findings "
                      "violate an allocated budget or environment; minor findings are documentation inconsistencies.")
    doc.add_picture(str(chart), width=Cm(17))
    doc.paragraphs[-1].alignment = WD_ALIGN_PARAGRAPH.CENTER


def heatmap(doc, findings, baseline):
    heading(doc, "2. Conflict heat map")
    doc.add_paragraph("Where the incoming document disagrees with the baseline: each cell counts the findings for one "
                      "category and one baseline document, coloured by the most severe of them.")
    docs = sorted({f["doc"] for f in findings})
    cats = sorted({f["category"] for f in findings})
    table = doc.add_table(rows=1, cols=len(docs) + 1)
    borders(table, color="FFFFFF", size=8)
    write(table.rows[0].cells[0], "Category", size=8, bold=True, color="FFFFFF")
    shade(table.rows[0].cells[0], NAVY)
    for cell, name in zip(table.rows[0].cells[1:], docs):
        shade(cell, NAVY)
        write(cell, name.replace("A3-", ""), size=7.5, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
    for cat in cats:
        cells = table.add_row().cells
        shade(cells[0], "EEF2F7")
        write(cells[0], cat, size=8.5, bold=True, color=NAVY, space=2)
        for cell, name in zip(cells[1:], docs):
            hits = [f for f in findings if f["category"] == cat and f["doc"] == name]
            if hits:
                worst = min(hits, key=lambda f: SEVERITIES.index(f["severity"]))["severity"]
                shade(cell, SEV_FILL[worst])
                write(cell, str(len(hits)), size=10, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
            else:
                shade(cell, "F3F4F6")
                write(cell, "·", size=9, color="9CA3AF", align=WD_ALIGN_PARAGRAPH.CENTER)
    first = 4.0
    widths(table, [first] + [(17 - first) / max(1, len(docs))] * len(docs))
    legend = doc.add_paragraph()
    legend.paragraph_format.space_before = Pt(4)
    for s in SEVERITIES:
        r = legend.add_run("■ ")
        r.font.color.rgb = RGBColor.from_string(SEV_FILL[s])
        r = legend.add_run(f"{s}    ")
        r.font.size = Pt(8)


def overview(doc, findings):
    heading(doc, "3. Findings at a glance")
    table = doc.add_table(rows=1, cols=5)
    borders(table)
    for cell, text in zip(table.rows[0].cells, ("ID", "Severity", "Finding", "Incoming req.", "Conflicts with")):
        shade(cell, NAVY)
        write(cell, text, size=8.5, bold=True, color="FFFFFF", space=2)
    for i, f in enumerate(findings):
        cells = table.add_row().cells
        write(cells[0], f["id"], size=8.5, bold=True, color=NAVY)
        shade(cells[1], SEV_FILL[f["severity"]])
        write(cells[1], f["severity"].upper(), size=7.5, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        write(cells[2], f["title"], size=8.5, space=2)
        write(cells[3], f["new_req"] or "—", size=8.5)
        write(cells[4], f"{f['baseline_req'] or '—'}\n{f['doc']}", size=8)
        if i % 2:
            for c in (cells[0], cells[2], cells[3], cells[4]):
                shade(c, "F8FAFC")
    widths(table, [1.4, 2.1, 7.1, 2.8, 3.6])


def details(doc, findings, new):
    heading(doc, "4. Detailed findings")
    for f in findings:
        heading(doc, f"{f['id']}  ·  {f['title']}", 2)
        tag = doc.add_table(rows=1, cols=2)
        borders(tag, size=0)
        shade(tag.rows[0].cells[0], SEV_FILL[f["severity"]])
        write(tag.rows[0].cells[0], f["severity"].upper(), size=8, bold=True, color="FFFFFF", align=WD_ALIGN_PARAGRAPH.CENTER)
        write(tag.rows[0].cells[1], f"   {f['category']}   ·   {f['new_req'] or '—'}  ↔  {f['baseline_req'] or '—'} ({f['doc']})",
              size=8.5, bold=True, color=GREY)
        widths(tag, [2.4, 14.6])
        compare = doc.add_table(rows=2, cols=2)
        borders(compare, color="CBD5E1")
        for cell, text, fill in ((compare.rows[0].cells[0], f"INCOMING  ·  {new['no']} Rev {new['rev']}  ·  {f['new_req']}", SEV_FILL[f["severity"]]),
                                 (compare.rows[0].cells[1], f"BASELINE  ·  {f['doc']}  ·  {f['baseline_req']}", "1E4E79")):
            shade(cell, fill)
            write(cell, text, size=8, bold=True, color="FFFFFF", space=2)
        for cell, text, tint in ((compare.rows[1].cells[0], f["new_text"], SEV_TINT[f["severity"]]),
                                 (compare.rows[1].cells[1], f["baseline_text"], "E8F0F8")):
            shade(cell, tint)
            write(cell, f"“{text}”" if text else "(no quote given)", size=9, italic=True, space=4)
        widths(compare, [8.5, 8.5])
        for label, text in (("Impact", f["impact"]), ("Recommended resolution", f["recommendation"])):
            if text:
                p = doc.add_paragraph()
                p.paragraph_format.space_before = Pt(5)
                r = p.add_run(label + ":  ")
                r.bold = True
                r.font.color.rgb = RGBColor.from_string(NAVY)
                p.add_run(text)


def actions(doc, findings):
    heading(doc, "5. Action plan")
    doc.add_paragraph("Proposed actions for the next Interface Working Group and Configuration Control Board.")
    table = doc.add_table(rows=1, cols=4)
    borders(table)
    for cell, text in zip(table.rows[0].cells, ("ID", "Action", "Owner", "Due")):
        shade(cell, NAVY)
        write(cell, text, size=8.5, bold=True, color="FFFFFF", space=2)
    for f in findings:
        owner = next((o for key, o in OWNERS.items() if key in f["category"].lower()), "Interceptor IPT")
        due = {"Critical": "Before PDR data package", "Major": "PDR", "Minor": "Next revision"}[f["severity"]]
        cells = table.add_row().cells
        write(cells[0], f["id"], size=8.5, bold=True, color=SEV_FILL[f["severity"]])
        write(cells[1], f["recommendation"] or f"Resolve: {f['title']}", size=8.5, space=2)
        write(cells[2], owner, size=8)
        write(cells[3], due, size=8)
    widths(table, [1.4, 9.8, 3.6, 2.2])


def appendix(doc, baseline, new):
    heading(doc, "Appendix A. Baseline documents checked")
    table = doc.add_table(rows=1, cols=4)
    borders(table)
    for cell, text in zip(table.rows[0].cells, ("Document", "Rev", "Title", "Requirements")):
        shade(cell, NAVY)
        write(cell, text, size=8.5, bold=True, color="FFFFFF", space=2)
    for d in baseline:
        cells = table.add_row().cells
        write(cells[0], d["no"], size=8.5, bold=True, color=NAVY)
        write(cells[1], d["rev"], size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
        write(cells[2], d["title"], size=8.5, space=2)
        write(cells[3], str(d["requirements"]), size=8.5, align=WD_ALIGN_PARAGRAPH.CENTER)
    widths(table, [3.8, 1.2, 9.6, 2.4])
    heading(doc, "Appendix B. How this report was produced")
    for step in (f"Trigger: {new['no']} Rev {new['rev']} was dropped onto the agent in EDK, waking both Hermes agents.",
                 "Baseline register (main Hermes): index/arrow3-baseline.md holds every requirement of every other "
                 "Word document uploaded so far; only new or changed documents are re-read.",
                 "Spec Reader (second Hermes, in parallel): converted the incoming document to Markdown with pandoc, "
                 "extracted every \"shall\" statement and wrote a management briefing.",
                 "Main Hermes studied the register and the new requirement lines together, then compared every "
                 "incoming requirement with the baseline.",
                 "This report was generated from the findings by the program report kit (consistency_report.py)."):
        doc.add_paragraph(step, style="List Number")
    p = doc.add_paragraph()
    r = p.add_run("Automated findings support, but do not replace, engineering review. Confirm each finding against "
                  "the source documents before CCB submission.")
    r.italic, r.font.size = True, Pt(8.5)
    r.font.color.rgb = RGBColor.from_string(GREY)


def main():
    parser = argparse.ArgumentParser(description=__doc__.splitlines()[0])
    parser.add_argument("findings", nargs="+", help="file(s) with the JSON findings")
    parser.add_argument("--index", default="index/arrow3-baseline.md")
    parser.add_argument("--source", default="", help="the incoming document")
    parser.add_argument("-o", "--output", required=True)
    args = parser.parse_args()
    raw = []
    for name in args.findings:
        raw += parse_findings(Path(name).read_text(errors="replace"))
    findings = normalise(raw)
    if not findings:
        sys.exit(f"consistency_report: no findings could be read from {', '.join(args.findings)}; "
                 "expected a JSON array of objects (id, severity, title, new_req, baseline_req, ...).")
    new = describe(args.source) if args.source else {"no": "Incoming document", "rev": "—", "title": ""}
    baseline = parse_index(Path(args.index))
    generated = dt.datetime.now()
    doc = Document()
    setup(doc, f"{new['no']} Rev {new['rev']}")
    with tempfile.TemporaryDirectory() as tmp:
        chart = charts(findings, Path(tmp))
        cover(doc, new, baseline, findings, generated)
        summary(doc, new, baseline, findings, chart)
        heatmap(doc, findings, baseline)
        overview(doc, findings)
        details(doc, findings, new)
        actions(doc, findings)
        appendix(doc, baseline, new)
        out = Path(args.output)
        out.parent.mkdir(parents=True, exist_ok=True)
        doc.core_properties.title = f"Requirements Consistency Report — {new['no']} Rev {new['rev']}"
        doc.core_properties.author = "EDK · Find Inconsistencies in Requirements"
        doc.save(out)
    print(f"Wrote {out}: {len(findings)} findings "
          f"({', '.join(f'{n} {s.lower()}' for s, n in Counter(f['severity'] for f in findings).items())}).")


if __name__ == "__main__":
    main()
